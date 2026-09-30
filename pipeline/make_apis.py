"""Convierte el relevamiento de agencias de promoción de inversiones (Excel, hoja "Resumen Apis") en apis.json.
Uso: python3 make_apis.py <archivo.xlsx> [salida.json]
Solo se publican datos institucionales de cada agencia (sin contactos personales).
"""
import json, os, sys, unicodedata, re
import openpyxl

def norm(x):
    x = unicodedata.normalize('NFD', str(x or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z]', '', x)

# País (texto del Excel, sin la aclaración entre paréntesis) -> ISO3, región
PAISES = {
    'argentina': 'ARG', 'china': 'CHN', 'suiza': 'CHE', 'noruega': 'NOR', 'eslovenia': 'SVN', 'alemania': 'DEU', 'portugal': 'PRT',
    'estadosunidos': 'USA', 'brasil': 'BRA', 'dinamarca': 'DNK', 'polonia': 'POL', 'paisesbajos': 'NLD', 'uruguay': 'URY',
    'canada': 'CAN', 'chile': 'CHL', 'francia': 'FRA', 'israel': 'ISR', 'finlandia': 'FIN', 'irlanda': 'IRL', 'espana': 'ESP',
    'japon': 'JPN', 'coreadelsur': 'KOR', 'australia': 'AUS', 'luxemburgo': 'LUX', 'arabiasaudita': 'SAU', 'peru': 'PER',
    'republicadominicana': 'DOM', 'sudafrica': 'ZAF', 'turkiye': 'TUR', 'turquia': 'TUR', 'grecia': 'GRC', 'nigeria': 'NGA',
    'kuwait': 'KWT', 'kiribati': 'KIR', 'macau': 'MAC', 'macao': 'MAC', 'argelia': 'DZA', 'djibouti': 'DJI', 'yibuti': 'DJI',
    'zambia': 'ZMB', 'mexico': 'MEX', 'colombia': 'COL', 'bolivia': 'BOL', 'ecuador': 'ECU', 'venezuela': 'VEN', 'austria': 'AUT',
    'bulgaria': 'BGR', 'chipre': 'CYP', 'croacia': 'HRV', 'eslovaquia': 'SVK', 'estonia': 'EST', 'hungria': 'HUN', 'letonia': 'LVA',
    'lituania': 'LTU', 'malta': 'MLT', 'republicacheca': 'CZE', 'chequia': 'CZE', 'rumania': 'ROU', 'suecia': 'SWE', 'panama': 'PAN',
    'paraguay': 'PRY', 'singapur': 'SGP', 'reinounido': 'GBR', 'belgica': 'BEL', 'emiratosarabesunidos': 'ARE', 'italia': 'ITA',
    'hongkong': 'HKG', 'india': 'IND', 'indonesia': 'IDN', 'vietnam': 'VNM', 'tailandia': 'THA', 'malasia': 'MYS', 'filipinas': 'PHL',
    'nuevazelanda': 'NZL', 'egipto': 'EGY', 'marruecos': 'MAR', 'kenia': 'KEN', 'qatar': 'QAT', 'catar': 'QAT', 'costarica': 'CRI',
    'guatemala': 'GTM', 'elsalvador': 'SLV', 'honduras': 'HND', 'jamaica': 'JAM', 'islandia': 'ISL', 'serbia': 'SRB', 'ucrania': 'UKR',
}
REGION = {}
for r, isos in {
    'América Latina y el Caribe': 'ARG BRA URY CHL PER DOM MEX COL BOL ECU VEN PAN PRY CRI GTM SLV HND JAM',
    'América del Norte': 'USA CAN',
    'Europa': 'CHE NOR SVN DEU PRT DNK POL NLD FRA FIN IRL ESP LUX TUR GRC AUT BGR CYP HRV SVK EST HUN LVA LTU MLT CZE ROU SWE GBR BEL ITA ISL SRB UKR',
    'Asia y Pacífico': 'CHN HKG MAC JPN KOR AUS KIR SGP IND IDN VNM THA MYS PHL NZL',
    'Medio Oriente y África': 'ISR SAU ZAF NGA KWT DZA DJI ZMB ARE EGY MAR KEN QAT',
}.items():
    for i in isos.split(): REGION[i] = r

SIGLA_FIX = {'KATRO': 'KOTRA'}  # error de tipeo en el relevamiento
# columnas del Excel (fila 2 de encabezados) -> clave de servicio
SERV = [('F', 'expo'), ('J', 'fin'), ('K', 'sl'), ('L', 'plat'), ('M', 'cap'), ('N', 'asis'), ('O', 'db'), ('P', 'inf'), ('Q', 'red'), ('S', 'news')]

def sino(v):
    t = norm(v)
    return True if t in ('si', 's') else False if t in ('no', 'n') else None

def clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v not in (None, '') else None

def num(v):
    try: return int(float(str(v).strip()))
    except Exception: return None

def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Resumen Apis'] if 'Resumen Apis' in wb.sheetnames else wb.worksheets[0]
    col = lambda L: openpyxl.utils.column_index_from_string(L) - 1
    ag, avisos = [], []
    for row in ws.iter_rows(min_row=4, values_only=True):
        sig, nom, pais = clean(row[col('B')]), clean(row[col('C')]), clean(row[col('D')])
        if not sig and not nom:
            continue
        base = re.sub(r'\(.*?\)', '', pais or '')
        m = re.search(r'\((.*?)\)', pais or '')
        iso = PAISES.get(norm(base))
        if norm(pais or '').startswith('chinahongkong'): iso = 'HKG'
        if not iso: avisos.append(f'País sin código: {pais} ({sig})')
        org = clean(row[col('G')])
        org = {'publica': 'Pública', 'mixta': 'Mixta', 'privada': 'Privada'}.get(norm(org), org)
        web = clean(row[col('T')])
        if web and not web.startswith('http'): web = 'https://' + web
        a = {'sigla': SIGLA_FIX.get(sig, sig), 'nombre': nom, 'pais': pais, 'subnacional': m.group(1) if m else None, 'iso': iso,
             'region': REGION.get(iso), 'depende': clean(row[col('E')]), 'org': org,
             'of_pais': num(row[col('H')]), 'of_ext': num(row[col('I')]), 'waipa': sino(row[col('R')]), 'web': web, 's': {}}
        for L, k in SERV:
            a['s'][k] = sino(row[col(L)])
        a['s']['ext'] = (a['of_ext'] or 0) > 0 if a['of_ext'] is not None else None
        a['promarg'] = norm(sig) == 'promargentina'
        ag.append(a)
    m8 = re.search(r'(\d{2})(\d{2})(\d{4})', os.path.basename(src))
    fecha = f'{m8.group(3)}-{m8.group(2)}-{m8.group(1)}' if m8 else None
    json.dump({'fuente': 'Relevamiento de agencias de promoción de inversiones (PromArgentina)', 'fecha': fecha,
               'agencias': ag}, open(out, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'ok {len(ag)} agencias, {len({a["iso"] for a in ag if a["iso"]})} países')
    for x in avisos: print('AVISO', x)

if __name__ == '__main__':
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'apis.json')
    main(src, out)
