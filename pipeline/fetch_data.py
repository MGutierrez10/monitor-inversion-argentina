"""Descarga datos oficiales para el Monitor Inversión Argentina y los guarda en data/.
Corre en GitHub Actions (tiene acceso directo a internet). No necesita claves.

Salidas:
  data/ied_cubo.json   IED por país x sector x subsector (BCRA), solo si el Excel cambió
  data/ied_meta.json   hash y último trimestre del Excel del BCRA
  data/series.json     series mensuales: IPC, tipo de cambio, reservas, inflación BCRA, riesgo país
  data/busqueda.json   candidatos de datos.gob.ar para EMAE y comercio exterior (para revisar IDs)
  data/estado.json     resultado de cada fuente (ok / error) y fecha de corrida
"""
import json, os, sys, hashlib, subprocess, datetime, urllib.request, urllib.parse, ssl, traceback

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
os.makedirs(DATA, exist_ok=True)
UA = {'User-Agent': 'Mozilla/5.0 (monitor-inversion-argentina; GitHub Actions)'}
estado = {'corrida': datetime.datetime.utcnow().strftime('%Y-%m-%dT%H:%MZ'), 'fuentes': {}}


def get(url, binary=False, timeout=90, insecure=False):
    ctx = ssl._create_unverified_context() if insecure else None
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as r:
        b = r.read()
    return b if binary else b.decode('utf-8', 'replace')


def load(name, default=None):
    p = os.path.join(DATA, name)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else default


def save(name, obj):
    json.dump(obj, open(os.path.join(DATA, name), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))


def fuente(nombre):
    def deco(fn):
        try:
            info = fn() or {}
            estado['fuentes'][nombre] = {'ok': True, **info}
        except Exception as e:
            estado['fuentes'][nombre] = {'ok': False, 'error': f'{type(e).__name__}: {e}'[:400]}
            traceback.print_exc()
        return fn
    return deco


# ---------- IED (BCRA, Anexo del Informe de IED)
IED_URLS = ['https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/anexo-informe-inversion-extranjera-directa.xlsx']

@fuente('ied_bcra')
def _ied():
    last_err = None
    for u in IED_URLS:
        for insecure in (False, True):
            try:
                b = get(u, binary=True, timeout=180, insecure=insecure); break
            except Exception as e:
                last_err = e; b = None
        if b: break
    if not b:
        raise last_err
    h = hashlib.sha256(b).hexdigest()
    meta = load('ied_meta.json', {})
    if meta.get('sha256') == h and os.path.exists(os.path.join(DATA, 'ied_cubo.json')):
        return {'cambio': False, 'ultimo_trimestre': meta.get('ultimo_trimestre')}
    tmp = os.path.join(DATA, '_bcra_ied.xlsx')
    open(tmp, 'wb').write(b)
    import openpyxl
    wb = openpyxl.load_workbook(tmp, read_only=True)
    faltan = [s for s in ('Tabla de datos I', 'Tabla de datos II', 'Tabla de datos IV', 'Tabla de datos V') if s not in wb.sheetnames]
    if faltan:
        os.remove(tmp)
        raise RuntimeError(f'El Excel del BCRA cambió de estructura: faltan hojas {faltan}. Hojas: {wb.sheetnames}')
    here = os.path.dirname(os.path.abspath(__file__))
    subprocess.run([sys.executable, os.path.join(here, 'make_ied.py'), tmp], check=True, cwd=here)
    os.replace(os.path.join(here, 'ied_cubo.json'), os.path.join(DATA, 'ied_cubo.json'))
    os.remove(tmp)
    cu = load('ied_cubo.json')
    ult = cu['trimestres'][-1]
    save('ied_meta.json', {'sha256': h, 'ultimo_trimestre': ult, 'descargado': estado['corrida'], 'url': u})
    return {'cambio': True, 'ultimo_trimestre': ult}


series = load('series.json', {}) or {}

# ---------- BCRA API v4.0 (reservas, tipo de cambio, inflación)
BCRA = {'reservas': 1, 'tc_mayorista': 5, 'ipc_mensual_bcra': 27, 'ipc_interanual_bcra': 28}

@fuente('bcra_api')
def _bcra():
    out = {}
    hoy = datetime.date.today().isoformat()
    for k, v in BCRA.items():
        u = f'https://api.bcra.gob.ar/estadisticas/v4.0/monetarias/{v}?desde=2017-01-01&hasta={hoy}&limit=3000'
        try:
            j = json.loads(get(u))
        except Exception:
            j = json.loads(get(u, insecure=True))
        res = j.get('results') or j.get('resultados') or []
        if isinstance(res, list) and res and 'detalle' in res[0]:
            res = res[0]['detalle']
        pts = sorted((r['fecha'][:10], float(r['valor'])) for r in res)
        series[k] = {'fuente': f'BCRA API v4.0, variable {v}', 'datos': pts}
        out[k] = pts[-1] if pts else None
    # agregados mensuales para el monitor
    def by_month(pts, how):
        m = {}
        for d, val in pts:
            m.setdefault(d[:7], []).append((d, val))
        return {k: (round(sum(x[1] for x in v) / len(v), 4) if how == 'avg' else sorted(v)[-1][1]) for k, v in sorted(m.items())}
    series['tc_prom_mensual'] = by_month(series['tc_mayorista']['datos'], 'avg')
    series['reservas_fin_mes'] = by_month(series['reservas']['datos'], 'last')
    return {'ultimos': out}


# ---------- datos.gob.ar (IPC INDEC)
SERIES_API = 'https://apis.datos.gob.ar/series/api/series/'

@fuente('ipc_indec')
def _ipc():
    q = urllib.parse.urlencode({'ids': '148.3_INIVELNAL_DICI_M_26:percent_change,148.3_INIVELNAL_DICI_M_26:percent_change_a_year_ago', 'start_date': '2016-12-01', 'format': 'json', 'limit': 1000})
    j = json.loads(get(SERIES_API + '?' + q))
    series['ipc'] = {'fuente': 'INDEC vía datos.gob.ar (148.3_INIVELNAL_DICI_M_26)', 'datos': j['data']}
    return {'ultimo': j['data'][-1]}


@fuente('busqueda_ids')
def _buscar():
    out = {}
    for q in ['estimador mensual actividad economica desestacionalizado', 'estimador mensual actividad economica original',
              'intercambio comercial exportaciones total mensual', 'intercambio comercial importaciones total mensual',
              'tipo de cambio real multilateral', 'producto interno bruto precios corrientes trimestral']:
        u = 'https://apis.datos.gob.ar/series/api/search/?' + urllib.parse.urlencode({'q': q, 'limit': 8})
        j = json.loads(get(u))
        out[q] = [{'id': r['field']['id'], 'titulo': r['field'].get('description') or r['field'].get('title'),
                   'frecuencia': r['field'].get('frequency'), 'hasta': r['field'].get('time_index_end')} for r in j.get('data', [])]
    save('busqueda.json', out)
    return {'consultas': len(out)}


# ---------- Riesgo país (ArgentinaDatos, serie histórica diaria del EMBI)
@fuente('riesgo_pais')
def _rp():
    j = json.loads(get('https://api.argentinadatos.com/v1/finanzas/indices/riesgo-pais'))
    pts = sorted((r['fecha'][:10], float(r['valor'])) for r in j if r.get('valor') is not None)
    series['riesgo_pais'] = {'fuente': 'J.P. Morgan EMBI vía ArgentinaDatos', 'ultimo': pts[-1],
                             'fin_mes': {m: v for m, v in {d[:7]: val for d, val in pts}.items()}}
    return {'ultimo': pts[-1]}


save('series.json', series)
save('estado.json', estado)
print(json.dumps(estado, ensure_ascii=False, indent=1))
